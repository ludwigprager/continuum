#!/usr/bin/env python3
"""
make_deck_template.py - builds templates/deck.potx. SPEC 6.4, 12.2.

    ./shell.sh python3 tools/make_deck_template.py --out templates/deck.potx

**This is not part of the daily pipeline, and it is a placeholder.** SPEC 12.2
is still open: the corporate PowerPoint template has not been supplied. Until
it arrives the deck is built against the plausible one this script writes, and
the moment the real `.potx` lands it replaces the output of this file - nothing
else changes but the mapping dict at the top of tools/render/pptx.py.

Why generate it rather than commit a hand-made file: the same reason
templates/workbook.xlsx is generated (SPEC 10, M3). A binary nobody can rebuild
is a binary nobody can review, and when the layouts need to move, whoever has
to move them needs PowerPoint. This keeps the placeholder in git as code.

What it writes:

  * 16:9, which is what a deck is now. python-pptx's own default template is
    4:3, so every placeholder is scaled horizontally to the wider slide -
    changing the slide size alone leaves the layouts sized for the old one.
  * the standard layout set, renamed to what the deck actually uses them for,
    so a reader of pptx.py can match LAYOUTS to what they see in PowerPoint.
  * Arial as the body font, the one metric-compatible choice that survives a
    trip through a corporate machine (SPEC 6.4 asks for Arial or Calibri).
  * a genuine `.potx`: the OPC content type of a PowerPoint *template*, not a
    presentation renamed. That matters - see the note in tools/render/pptx.py
    about python-pptx refusing to open one. Writing a real template here is
    what proves the loader handles the file the corporate template will be.

Exit codes: 0 ok, 2 tool/usage error.
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
import zipfile
from pathlib import Path
from typing import Sequence

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:  # pragma: no cover
    sys.exit("python-pptx is required (it is in the pipeline image)")

EXIT_OK, EXIT_TOOL = 0, 2

DEFAULT_OUT = Path("templates/deck.potx")

# 16:9 at the size PowerPoint itself uses. Height is unchanged from the 4:3
# default, so only the horizontal axis is rescaled below.
SLIDE_WIDTH = Emu(12192000)
SLIDE_HEIGHT = Emu(6858000)

BODY_FONT = "Arial"

# The layout the chart slides use, repositioned by `relayout` below. It is the
# same index tools/render/pptx.py names in its LAYOUTS dict.
PICTURE_LAYOUT = 8

# An explicit position, as opposed to one inherited from the slide master.
XFRM = "{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm"

# The standard layouts, renamed to the job the deck gives them. The index is
# what pptx.py's mapping refers to; the name is what a human sees in
# PowerPoint, and pptx.py checks the two still agree.
LAYOUT_NAMES = {
    0: "Title Slide",
    1: "Title and Content",
    2: "Section Header",
    3: "Two Content",
    5: "Title Only",
    6: "Blank",
    8: "Picture with Caption",
}

# Every <a:latin typeface="..."> in the theme - major and minor font alike.
TYPEFACE = re.compile(rb'(<a:latin typeface=")[^"]*"')

# The OPC content type of a template, as opposed to a presentation. A .potx
# that carries the presentation type is a .pptx with the wrong extension.
PRESENTATION_TYPE = ("application/vnd.openxmlformats-officedocument"
                     ".presentationml.presentation.main+xml")
TEMPLATE_TYPE = ("application/vnd.openxmlformats-officedocument"
                 ".presentationml.template.main+xml")

# A fixed timestamp for every zip entry and for the document properties, so
# the committed template does not change every time it is regenerated
# (SPEC 5.2, same rule as the workbook).
ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)
EPOCH = dt.datetime(1980, 1, 1)


def has_own_geometry(shape) -> bool:
    """Does this shape carry its own position, or inherit one?

    A layout placeholder usually has no `<a:xfrm>` of its own: it takes its
    box from the same placeholder on the master. Reading `shape.left` hides
    that - python-pptx resolves the inheritance and hands back the master's
    value - but *writing* it does not: it creates an `<a:off>` with no
    matching `<a:ext>`, and the placeholder ends up at height zero, which is
    how a title comes to be a sliver at the top edge of the slide.
    """
    return shape._element.find(f".//{XFRM}") is not None


def widen(presentation) -> None:
    """Rescale the 4:3 default to 16:9.

    Setting `slide_width` alone makes the slide wider and leaves everything on
    it laid out for the narrower one - a title that stops two thirds of the
    way across and a picture frame that does not fill the space. The
    horizontal axis is the only one that moves, so left and width scale and
    top and height do not. Only shapes that own their geometry are touched;
    the ones that inherit it follow the master for free.
    """
    scale = SLIDE_WIDTH / presentation.slide_width
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT

    for master in presentation.slide_masters:
        shapes = list(master.shapes) + [shape for layout in master.slide_layouts
                                        for shape in layout.shapes]
        for shape in shapes:
            if not has_own_geometry(shape):
                continue
            shape.left = Emu(int(shape.left * scale))
            shape.width = Emu(int(shape.width * scale))


def relayout(presentation) -> None:
    """Give the chart layout a shape a report can use.

    python-pptx's "Picture with Caption" puts the picture first and the title
    underneath it, which reads as a caption above a caption. A deck slide
    wants the title at the top, the chart in the middle and the counting rule
    under it, so that layout is positioned explicitly here - in the template,
    where a design decision belongs, rather than in the renderer, which only
    fills what it is given (SPEC 11).
    """
    width, height = presentation.slide_width, presentation.slide_height
    boxes = {
        # idx: (left, top, width, height) as fractions of the slide
        0: (0.06, 0.05, 0.88, 0.11),   # title
        1: (0.06, 0.18, 0.88, 0.66),   # picture
        2: (0.06, 0.86, 0.88, 0.09),   # caption
    }
    layout = presentation.slide_layouts[PICTURE_LAYOUT]
    for shape in layout.placeholders:
        box = boxes.get(shape.placeholder_format.idx)
        if box is None:
            continue
        left, top, box_width, box_height = box
        shape.left = Emu(int(width * left))
        shape.top = Emu(int(height * top))
        shape.width = Emu(int(width * box_width))
        shape.height = Emu(int(height * box_height))


def rename_layouts(presentation, names: dict[int, str]) -> list[str]:
    """Give the layouts the names pptx.py expects, and report the result."""
    layouts = presentation.slide_layouts
    for index, name in names.items():
        if index < len(layouts):
            layouts[index].name = name
    return [layout.name for layout in layouts]


def to_template(path: Path, font: str) -> int:
    """Rewrite the saved .pptx as a genuine .potx, in place.

    Three things happen here rather than through python-pptx, because
    python-pptx models none of them:

      * the OPC content type. A template differs from a presentation by that
        one string; python-pptx can only write the presentation one.
      * the theme fonts. python-pptx does not model the theme, so it hands out
        a generic part with bytes and no element tree - setting `typeface` on
        it does nothing at all, quietly. SPEC 6.4 asks for Arial or Calibri
        and a font nobody checked is a font that gets substituted silently
        (SPEC 8.2), so the substitution is counted and reported.
      * every entry's timestamp, pinned, so regenerating the template twice
        gives the same bytes and git sees a change only when one happened.
    """
    replaced = 0
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename))
                   for info in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, payload in entries:
            if info.filename == "[Content_Types].xml":
                payload = payload.replace(PRESENTATION_TYPE.encode(),
                                          TEMPLATE_TYPE.encode())
            if info.filename.startswith("ppt/theme/"):
                payload, count = TYPEFACE.subn(
                    rb'\g<1>' + font.encode() + rb'"', payload)
                replaced += count
            copy = zipfile.ZipInfo(info.filename, date_time=ZIP_EPOCH)
            copy.compress_type = info.compress_type
            copy.external_attr = info.external_attr
            copy.internal_attr = info.internal_attr
            copy.create_system = info.create_system
            archive.writestr(copy, payload)
    return replaced


def build(out: Path, font: str = BODY_FONT) -> tuple[list[str], int]:
    presentation = Presentation()  # python-pptx's own default, then reshaped
    widen(presentation)
    relayout(presentation)
    names = rename_layouts(presentation, LAYOUT_NAMES)

    # All of it set explicitly. python-pptx's default template carries its
    # own author and dates, and left alone they end up in a file this project
    # ships - a corporate deck template crediting a stranger, stamped with a
    # date nobody here chose.
    properties = presentation.core_properties
    properties.title = "continuum deck template (placeholder)"
    properties.author = "continuum pipeline"
    properties.last_modified_by = "continuum pipeline"
    properties.comments = (
        "Placeholder until the corporate template is supplied (SPEC 12.2). "
        "Regenerate with tools/make_deck_template.py.")
    properties.created = EPOCH
    properties.modified = EPOCH
    properties.revision = 1

    out.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(out)
    fonts = to_template(out, font)
    if not fonts:
        raise RuntimeError(
            f"the theme carried no <a:latin> typeface, so {font} was not set "
            f"anywhere. The deck would be laid out in whatever PowerPoint "
            f"substitutes (SPEC 6.4, 8.2).")
    return names, fonts


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Build the placeholder templates/deck.potx (SPEC 12.2).")
    parser.add_argument("--out", default=DEFAULT_OUT, type=Path)
    parser.add_argument("--font", default=BODY_FONT,
                        help="body font (SPEC 6.4: Arial or Calibri)")
    args = parser.parse_args(argv)

    try:
        names, fonts = build(args.out, args.font)
    except Exception as exc:  # pragma: no cover - surfaced as a tool error
        print(f"make_deck_template: {type(exc).__name__}: {exc}",
              file=sys.stderr)
        return EXIT_TOOL

    print(f"deck template: {args.out}")
    print(f"  {SLIDE_WIDTH} x {SLIDE_HEIGHT} EMU, "
          f"{args.font} set on {fonts} theme font slots")
    for index, name in enumerate(names):
        print(f"  layout {index:>2}  {name}")
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())

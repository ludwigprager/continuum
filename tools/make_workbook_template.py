#!/usr/bin/env python3
"""
make_workbook_template.py - builds templates/workbook.xlsx. SPEC 6.4.

    ./shell.sh python3 tools/make_workbook_template.py \\
        --model out/reports/<date>/report_model.json \\
        --spec reports/daily.yaml --out templates/workbook.xlsx

**This is not part of the daily pipeline.** It is run when the pre-built
pivots change, or when a schema change moves the columns a pivot reads - which
tools/render/xlsx.py warns about on every run, naming this command.

Why it exists: openpyxl cannot create a pivot table from nothing, so SPEC 6.4
says ship a template with the pivots already in it, load it, and write the data
into it. Something has to write that template once. Doing it in code rather
than by hand in Excel keeps it in git, diffable and reproducible, and means a
junior can regenerate it after a schema change without owning a copy of Excel.

Which pivots exist is a report decision, so it lives in reports/daily.yaml
under `pivots:` next to the tables, not here. This file contains no field name
(SPEC 6.1.1).

What the template holds:

  * one sheet per grain that a pivot reads, with the header row and an empty
    data row, wrapped in a real Excel table. The pivot cache points at the
    table by name, so it follows the data however many rows arrive.
  * the pivot tables themselves, with empty caches marked refreshOnLoad, so
    Excel fills them in from the data xlsx.py wrote.

Exit codes: 0 ok, 2 tool/usage error.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Sequence

try:
    from openpyxl import Workbook
    from openpyxl.pivot.cache import (CacheDefinition, CacheField, CacheSource,
                                      SharedItems, WorksheetSource)
    from openpyxl.pivot.record import RecordList
    from openpyxl.pivot.table import (DataField, Location, PivotField,
                                      PivotTableStyle, RowColField,
                                      TableDefinition)
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required (it is in the pipeline image)")

try:
    from ruamel.yaml import YAML
except ImportError:  # pragma: no cover
    sys.exit("ruamel.yaml is required (it is in the pipeline image)")

sys.path.insert(0, str(Path(__file__).resolve().parent / "render"))
import xlsx  # noqa: E402  - the layout the template must match, not a copy of it

EXIT_OK, EXIT_TOOL = 0, 2

# Excel's own version markers. Without them Excel treats the pivot as one
# written by a very old producer and refuses to refresh it.
PIVOT_VERSION = 8
MIN_REFRESHABLE = 3


def resolve(columns: Sequence[dict], field: str) -> int:
    """Index of the sheet column a pivot field names.

    `rows: datacenter` in the spec means the column that holds the code, which
    is `datacenter_code` once build_model.py has marked it coded. Resolving it
    here is what keeps the enum vocabulary out of this file.
    """
    headers = [column["header"] for column in columns]
    for candidate in (f"{field}_code", field):
        if candidate in headers:
            return headers.index(candidate)
    raise KeyError(f"{field!r} is not a column of this grain: {headers}")


def build_pivot(workbook: Workbook, key: str, spec: dict,
                columns: Sequence[dict], lang: str) -> None:
    grain = spec["grain"]
    table_name = f"{xlsx.TABLE_PREFIX}{grain}"
    headers = [column["header"] for column in columns]

    row_field = resolve(columns, spec["rows"])
    col_field = resolve(columns, spec["columns"])
    measure = resolve(columns, spec["measure"])

    cache = CacheDefinition(
        cacheSource=CacheSource(
            type="worksheet",
            worksheetSource=WorksheetSource(name=table_name)),
        cacheFields=[CacheField(name=header, sharedItems=SharedItems())
                     for header in headers],
        # The cache ships empty: the data is written after the template, by
        # xlsx.py. Excel rebuilds it from the table on open.
        refreshOnLoad=True,
        recordCount=0,
        createdVersion=PIVOT_VERSION,
        refreshedVersion=PIVOT_VERSION,
        minRefreshableVersion=MIN_REFRESHABLE,
    )
    cache.records = RecordList()

    fields = []
    for index in range(len(headers)):
        if index == row_field:
            fields.append(PivotField(axis="axisRow", showAll=False,
                                     compact=False, outline=False))
        elif index == col_field:
            fields.append(PivotField(axis="axisCol", showAll=False,
                                     compact=False, outline=False))
        elif index == measure:
            fields.append(PivotField(dataField=True, showAll=False,
                                     compact=False, outline=False))
        else:
            fields.append(PivotField(showAll=False, compact=False, outline=False))

    sheet = workbook.create_sheet(spec.get("sheet", f"pivot_{key}"))
    xlsx.put(sheet, 1, 1, xlsx.t(spec, "title", lang, key), font=xlsx.TITLE)
    xlsx.put(sheet, 2, 1, f"{key}  |  {spec.get('definition', '')}  |  "
                          f"{grain}")

    pivot = TableDefinition(
        name=f"pivot_{key}",
        cacheId=1,
        dataOnRows=False,
        dataCaption=xlsx.t(spec, "measure_caption", lang, "Count"),
        location=Location(ref="A4:D16", firstHeaderRow=1, firstDataRow=2,
                          firstDataCol=1),
        pivotFields=fields,
        rowFields=[RowColField(x=row_field)],
        colFields=[RowColField(x=col_field)],
        dataFields=[DataField(name=xlsx.t(spec, "measure_caption", lang, "Count"),
                              fld=measure, subtotal="count", baseField=-1,
                              baseItem=0)],
        pivotTableStyleInfo=PivotTableStyle(name="PivotStyleLight16",
                                            showRowHeaders=True,
                                            showColHeaders=True,
                                            showLastColumn=True),
        updatedVersion=PIVOT_VERSION,
        createdVersion=PIVOT_VERSION,
        minRefreshableVersion=MIN_REFRESHABLE,
    )
    pivot.cache = cache
    sheet.add_pivot(pivot)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build templates/workbook.xlsx.")
    parser.add_argument("--model", required=True, type=Path,
                        help="a report_model.json - the template's columns "
                             "must be the ones the renderer writes")
    parser.add_argument("--spec", default=Path("reports/daily.yaml"), type=Path)
    parser.add_argument("--out", default=Path("templates/workbook.xlsx"), type=Path)
    parser.add_argument("--lang", default="de", choices=("de", "en"))
    args = parser.parse_args(argv)

    for required in (args.model, args.spec):
        if not required.is_file():
            print(f"{required} not found", file=sys.stderr)
            return EXIT_TOOL

    model = json.loads(args.model.read_text(encoding="utf-8"))
    spec = YAML(typ="safe").load(args.spec.read_text(encoding="utf-8")) or {}
    pivots = spec.get("pivots") or {}
    if not pivots:
        print(f"{args.spec}: no `pivots:` section, nothing to build",
              file=sys.stderr)
        return EXIT_TOOL

    workbook = Workbook()
    workbook.remove(workbook.active)

    layouts: dict[str, list[dict]] = {}
    for key in sorted(pivots):
        grain = pivots[key]["grain"]
        specs = (model.get("flat_columns") or {}).get(grain)
        if not specs:
            print(f"{args.spec}: pivot {key!r} reads {grain!r}, which the "
                  f"model does not have", file=sys.stderr)
            return EXIT_TOOL
        columns = xlsx.sheet_columns(specs, args.lang)
        if grain not in layouts:
            layouts[grain] = columns
            # The source sheet carries headers and one empty row: a pivot
            # cache has to name real columns, and an Excel table needs a body.
            # xlsx.py overwrites both with the real data.
            xlsx.write_data_sheet(workbook.create_sheet(grain), grain,
                                  columns, [])
        try:
            build_pivot(workbook, key, pivots[key], layouts[grain], args.lang)
        except KeyError as exc:
            print(f"{args.spec}: pivot {key!r}: {exc}", file=sys.stderr)
            return EXIT_TOOL

    args.out.parent.mkdir(parents=True, exist_ok=True)
    workbook.properties.creator = "continuum pipeline"
    # A fixed date, not the clock: the template is committed, and a rebuild
    # that changed nothing should produce no diff.
    workbook.properties.created = xlsx.EPOCH
    workbook.properties.modified = xlsx.EPOCH
    workbook.save(args.out)
    xlsx.normalise(args.out, xlsx.EPOCH)

    print(f"template: {args.out}")
    for key in sorted(pivots):
        print(f"  pivot {key} on {pivots[key]['grain']} "
              f"({pivots[key]['rows']} x {pivots[key]['columns']})")
    print("  the data sheets here are placeholders; xlsx.py writes the data")
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())

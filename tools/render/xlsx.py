#!/usr/bin/env python3
"""
xlsx.py - the Excel report. SPEC 6.4, milestone M3.

This is the output that actually gets used: managers open it and slice the
data themselves. Everything else is downstream of that.

    tools/render/xlsx.py --model report_model.json --out DIR [--lang de|en]

The renderer is dumb on purpose (SPEC 11). Every number, label and column
already exists in report_model.json; this file lays them out and computes
nothing. It never opens taxonomy.yaml - the labels were baked in by
build_model.py - and it contains no field name, no code and no enum value
(SPEC 6.1.1): the sheets are built from `flat_columns`, whatever that happens
to contain today.

Sheets:

    dashboard      headline numbers, coverage, provenance, the chart PNGs
    counts         every frequency table stacked vertically, no pivot needed
    projects       one row per project, plus the convenience columns
    datastores     one row per (project, datastore)
    environments   one row per (project, environment)
    blockers       one row per (project, blocker)
    dependencies   one row per edge
    definitions    the counting rules, one line each
    pivot_*        whatever pre-built pivots the template carries

Two columns per coded field: `security_class_code` holds S3 and
`security_class_label` holds the label. Pivot on the code so the sort order is
the taxonomy's, display the label.

**The template matters.** openpyxl cannot create a pivot table from nothing, so
templates/workbook.xlsx is loaded and written into rather than built from
scratch. Build from scratch and the pivots are gone - which is why a missing
template is a loud warning here and a failing test in tests/test_xlsx.py,
rather than a workbook that quietly lost a feature.

Exit codes: 0 ok, 2 tool/usage error.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any, Sequence

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required (it is in the pipeline image)")

# A renderer has no verdict on the data, so it never exits 1.
EXIT_OK, EXIT_TOOL = 0, 2

DEFAULT_TEMPLATE = Path("templates/workbook.xlsx")
OUTPUT_NAME = "report.xlsx"

# Grain sheets, in the order they appear in the workbook. Taken from the model
# rather than declared here would be nicer still, but the workbook needs a
# stable order and `flat` is a mapping.
DATA_SHEETS = ("projects", "datastores", "environments", "blockers",
               "dependencies")
SHEET_ORDER = ("dashboard", "counts", *DATA_SHEETS, "definitions")

TABLE_PREFIX = "tbl_"

# Arial or Calibri, nothing exotic (SPEC 6.4).
BODY_FONT = "Calibri"
HEAD = Font(name=BODY_FONT, bold=True)
PLAIN = Font(name=BODY_FONT)
TITLE = Font(name=BODY_FONT, bold=True, size=14)

# Anything heavier fights pivot tables.
TABLE_STYLE = TableStyleInfo(name="TableStyleLight1", showRowStripes=False,
                             showColumnStripes=False, showFirstColumn=False,
                             showLastColumn=False)

COUNT_FORMAT = "0"
PCT_FORMAT = "0.0"
MIN_WIDTH, MAX_WIDTH = 9, 44

# A fixed timestamp for every zip entry. Two runs of the same model then
# produce the same bytes, the way report_model.json does (SPEC 5.2).
ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)
EPOCH = dt.datetime(1980, 1, 1)


def stamp(generated_at: Any) -> dt.datetime:
    """The workbook's created/modified date, taken from the model.

    openpyxl defaults both to datetime.now(), which would make two runs of the
    same model differ. The model's own timestamp is the honest answer and is
    already pinned wherever reproducibility is being tested.
    """
    try:
        parsed = dt.datetime.fromisoformat(str(generated_at))
    except (TypeError, ValueError):
        return EPOCH
    return parsed.replace(tzinfo=None)


def t(item: dict, field: str, lang: str, default: Any = "") -> Any:
    """`field` in the requested language, falling back to the other one."""
    for candidate in (f"{field}_{lang}", f"{field}_de", f"{field}_en", field):
        value = item.get(candidate)
        if value not in (None, ""):
            return value
    return default


# --------------------------------------------------------------------------
# Cells
# --------------------------------------------------------------------------

def put(ws: Worksheet, row: int, column: int, value: Any,
        font: Font = PLAIN, number_format: str | None = None) -> None:
    """Write one cell, as text unless it is genuinely a number.

    A string is forced to text even when it starts with `=`. Project files are
    hand-edited by several hundred people; a cell that begins with an equals
    sign is a typo, not a formula the report should evaluate.
    """
    cell = ws.cell(row=row, column=column)
    if isinstance(value, bool):
        value = str(value).lower()
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = font
    if number_format:
        cell.number_format = number_format


def autosize(ws: Worksheet, widths: dict[int, int]) -> None:
    for index, width in widths.items():
        ws.column_dimensions[get_column_letter(index)].width = \
            max(MIN_WIDTH, min(MAX_WIDTH, width + 2))


def measure(widths: dict[int, int], column: int, value: Any) -> None:
    widths[column] = max(widths.get(column, 0), len(str(value if value is not None else "")))


# --------------------------------------------------------------------------
# The data sheets
# --------------------------------------------------------------------------

def sheet_columns(specs: Sequence[dict], lang: str) -> list[dict]:
    """Model columns -> sheet columns.

    A coded column becomes two: the code, which pivots and sorts in the
    taxonomy's order, and the label, which is what a human reads. Everything
    else becomes one. The headers do not depend on the language, so a template
    built for one language fits the other (SPEC 6.4).
    """
    columns: list[dict] = []
    for spec in specs:
        key = spec["key"]
        if spec.get("coded"):
            columns.append({"header": f"{key}_code", "key": key, "kind": "code"})
            columns.append({"header": f"{key}_label",
                            "key": f"{key}_label_{lang}", "kind": "text"})
        else:
            columns.append({"header": key, "key": key,
                            "kind": spec.get("kind", "text")})
    return columns


def write_data_sheet(ws: Worksheet, name: str, columns: Sequence[dict],
                     rows: Sequence[dict]) -> None:
    """One grain, one real Excel table: autofilter, frozen header, no styling
    to speak of. This is what gets pivoted."""
    widths: dict[int, int] = {}
    for index, column in enumerate(columns, start=1):
        put(ws, 1, index, column["header"], font=HEAD)
        measure(widths, index, column["header"])

    for offset, row in enumerate(rows, start=2):
        for index, column in enumerate(columns, start=1):
            value = row.get(column["key"])
            fmt = COUNT_FORMAT if column["kind"] == "count" else None
            put(ws, offset, index, value, number_format=fmt)
            measure(widths, index, value)

    autosize(ws, widths)
    ws.freeze_panes = "A2"

    # An Excel table needs at least one data row in its range, even when there
    # is no data yet: the first day nobody has recorded a blocker, the sheet
    # must still be filterable and still be a valid pivot source.
    last_row = max(2, len(rows) + 1)
    ref = f"A1:{get_column_letter(len(columns))}{last_row}"
    table = Table(displayName=f"{TABLE_PREFIX}{name}", ref=ref)
    table.tableStyleInfo = TABLE_STYLE
    ws.add_table(table)


# --------------------------------------------------------------------------
# counts: every frequency table stacked vertically
# --------------------------------------------------------------------------

def write_counts_sheet(ws: Worksheet, model: dict, lang: str) -> int:
    """The sheet that answers most of what gets asked.

    No pivot skills required, and it is what people paste into emails. A table
    the data cannot support yet is printed with its note rather than skipped -
    an absent row and a zero row mean different things (SPEC 11).
    """
    widths: dict[int, int] = {}
    row = 1
    put(ws, row, 1, t(model.get("report", {}), "title", lang, "Report"), font=TITLE)
    row += 1
    put(ws, row, 1, f"as of {model.get('as_of_date', '')}")
    row += 2

    for key in sorted(model.get("tables") or {}):
        table = model["tables"][key]
        put(ws, row, 1, t(table, "title", lang, key), font=HEAD)
        measure(widths, 1, t(table, "title", lang, key))
        row += 1
        put(ws, row, 1, f"{key}  |  {table.get('definition', '')}  |  "
                        f"{table.get('grain', '')}/{table.get('counts', '')}")
        row += 1

        if not table.get("available", False):
            put(ws, row, 1, t(table, "note", lang, "n/a"))
            row += 2
            continue

        columns = table.get("columns") or []
        rows = table.get("rows") or []
        first = columns[0]["key"] if columns else None

        if table.get("kind") == "cross_tab" or len(columns) > 2:
            # A cross-tab keeps its shape: row label, one column per code,
            # then the row total.
            put(ws, row, 1, t(columns[0], "label", lang, first), font=HEAD)
            for index, column in enumerate(columns[1:], start=2):
                header = t(column, "label", lang, column["key"])
                put(ws, row, index, header, font=HEAD)
                measure(widths, index, header)
            row += 1
            for record in rows:
                label = record.get(f"{first}_label_{lang}", record.get(first))
                put(ws, row, 1, f"{record.get(first)}  {label}")
                measure(widths, 1, f"{record.get(first)}  {label}")
                for index, column in enumerate(columns[1:], start=2):
                    put(ws, row, index, record.get(column["key"], 0),
                        number_format=COUNT_FORMAT)
                row += 1
        else:
            for index, header in enumerate(("Code", "Label", "Anzahl / Count"),
                                           start=1):
                put(ws, row, index, header, font=HEAD)
                measure(widths, index, header)
            row += 1
            for record in rows:
                code = record.get(first)
                put(ws, row, 1, code)
                measure(widths, 1, code)
                label = record.get(f"{first}_label_{lang}", code)
                put(ws, row, 2, label)
                measure(widths, 2, label)
                put(ws, row, 3, record.get("value", 0), number_format=COUNT_FORMAT)
                row += 1
        row += 1

    autosize(ws, widths)
    ws.freeze_panes = "A4"
    return row


# --------------------------------------------------------------------------
# definitions: the counting rules, printed next to the numbers
# --------------------------------------------------------------------------

def write_definitions_sheet(ws: Worksheet, model: dict, lang: str) -> None:
    """SPEC 7: every rule is printed on a sheet in the Excel, so a number and
    the rule it was counted under never travel separately."""
    headers = ("key", "grain", "counts", "definition")
    for index, header in enumerate(headers, start=1):
        put(ws, 1, index, header, font=HEAD)

    for offset, definition in enumerate(model.get("definitions") or [], start=2):
        put(ws, offset, 1, definition.get("key"))
        put(ws, offset, 2, definition.get("grain"))
        put(ws, offset, 3, definition.get("counts"))
        put(ws, offset, 4, t(definition, "text", lang))
        ws.cell(row=offset, column=4).alignment = Alignment(wrap_text=True,
                                                            vertical="top")

    autosize(ws, {1: 30, 2: 14, 3: 10})
    ws.column_dimensions["D"].width = 110
    ws.freeze_panes = "A2"


# --------------------------------------------------------------------------
# dashboard: the numbers, the provenance, and the charts
# --------------------------------------------------------------------------

def write_dashboard_sheet(ws: Worksheet, model: dict, lang: str,
                          model_dir: Path) -> list[str]:
    """Returns the chart keys that could not be embedded, for the caller to
    report. charts.py (M4) writes the PNGs; until it exists the sheet says so
    rather than pretending they were never asked for."""
    widths: dict[int, int] = {}
    row = 1
    put(ws, row, 1, t(model.get("report", {}), "title", lang, "Report"), font=TITLE)
    row += 2

    coverage = model.get("coverage") or {}
    provenance = model.get("provenance") or {}
    facts = [
        ("as_of_date", model.get("as_of_date")),
        ("generated_at", model.get("generated_at")),
        ("git_sha", provenance.get("git_sha")),
        ("git_tag", provenance.get("git_tag")),
        ("schema_version", provenance.get("schema_version")),
        ("pipeline_version", provenance.get("pipeline_version")),
        ("image_digest", provenance.get("image_digest")),
    ]
    for key, value in facts:
        put(ws, row, 1, key, font=HEAD)
        put(ws, row, 2, value)
        measure(widths, 1, key)
        measure(widths, 2, value)
        row += 1
    row += 1

    for item in model.get("headline") or []:
        label = t(item, "label", lang, item.get("key"))
        put(ws, row, 1, label, font=HEAD)
        measure(widths, 1, label)
        value = item.get("value")
        put(ws, row, 2, "n/a" if value is None else value,
            number_format=None if value is None else COUNT_FORMAT)
        put(ws, row, 3, item.get("definition"))
        measure(widths, 3, item.get("definition"))
        row += 1
    row += 1

    # Raw coverage and verified coverage, never merged into one number
    # (SPEC 11). An import gives 90% raw coverage that nobody has looked at.
    put(ws, row, 1, "coverage", font=HEAD)
    row += 1
    for key, fmt in (("projects_total", COUNT_FORMAT),
                     ("fields_tracked", COUNT_FORMAT),
                     ("field_coverage_pct", PCT_FORMAT),
                     ("verified_coverage_pct", PCT_FORMAT)):
        put(ws, row, 1, key)
        measure(widths, 1, key)
        put(ws, row, 2, coverage.get(key), number_format=fmt)
        row += 1
    row += 1

    team_headers = ("team_id", "projects", "coverage_pct", "verified_pct")
    for index, header in enumerate(team_headers, start=1):
        put(ws, row, index, header, font=HEAD)
        measure(widths, index, header)
    row += 1
    for team in coverage.get("by_team") or []:
        put(ws, row, 1, team.get("team_id"))
        measure(widths, 1, team.get("team_id"))
        put(ws, row, 2, team.get("projects"), number_format=COUNT_FORMAT)
        put(ws, row, 3, team.get("coverage_pct"), number_format=PCT_FORMAT)
        put(ws, row, 4, team.get("verified_pct"), number_format=PCT_FORMAT)
        row += 1
    row += 1

    autosize(ws, widths)

    missing: list[str] = []
    for chart in model.get("charts") or []:
        put(ws, row, 1, t(chart, "title", lang, chart.get("key")), font=HEAD)
        row += 1
        path = model_dir / str(chart.get("path", ""))
        if path.is_file():
            row = embed(ws, path, row, chart)
        else:
            missing.append(str(chart.get("key")))
            put(ws, row, 1, f"[{path.name} not rendered yet]")
            row += 2
    return missing


def embed(ws: Worksheet, path: Path, row: int, chart: dict) -> int:
    from openpyxl.drawing.image import Image  # needs Pillow

    image = Image(str(path))
    # The model declares the size the chart was rendered at; scale it to
    # something that fits a screen without touching the aspect ratio.
    width = int(chart.get("width_px") or image.width)
    height = int(chart.get("height_px") or image.height)
    scale = 640 / width if width else 1
    image.width, image.height = int(width * scale), int(height * scale)
    ws.add_image(image, f"A{row}")
    return row + int(image.height / 19) + 2


# --------------------------------------------------------------------------
# Template
# --------------------------------------------------------------------------

def open_workbook(template: Path | None) -> tuple[Workbook, list[str]]:
    """Load the template, or say loudly that there is no template.

    openpyxl cannot create a pivot table (SPEC 6.4). The pre-built pivots live
    in templates/workbook.xlsx and survive because the workbook is loaded and
    written into. Building one from scratch loses them, so that path warns
    instead of quietly shipping a workbook with a feature missing.
    """
    warnings: list[str] = []
    if template is None or not template.is_file():
        warnings.append(
            f"no template at {template}: the pre-built pivot tables are NOT in "
            f"this workbook. Regenerate it with "
            f"tools/make_workbook_template.py (SPEC 6.4).")
        workbook = Workbook()
        workbook.remove(workbook.active)
        return workbook, warnings
    return load_workbook(template), warnings


def sheet(workbook: Workbook, name: str) -> Worksheet:
    """An empty sheet by that name, keeping the template's own if it had one.

    A template sheet is emptied rather than replaced: deleting and recreating
    it would drop the pivot cache's reference to the table it reads.
    """
    if name in workbook.sheetnames:
        existing = workbook[name]
        for table in list(existing.tables):
            del existing.tables[table]
        existing.delete_rows(1, existing.max_row + 1)
        return existing
    return workbook.create_sheet(name)


def check_template(workbook: Workbook, layouts: dict[str, list[dict]]) -> list[str]:
    """Do the template's pivot caches still match the columns being written?

    A pivot reads its source by field name. Add a field to the schema and the
    columns move; the pivot then counts the wrong one or breaks on open. That
    must be a message, not a surprise in a management meeting.
    """
    warnings: list[str] = []
    for name in workbook.sheetnames:
        for pivot in workbook[name]._pivots:
            source = getattr(pivot.cache.cacheSource, "worksheetSource", None)
            table = getattr(source, "name", None) or ""
            grain = table[len(TABLE_PREFIX):] if table.startswith(TABLE_PREFIX) else ""
            if grain not in layouts:
                continue
            cached = [field.name for field in pivot.cache.cacheFields]
            written = [column["header"] for column in layouts[grain]]
            if cached != written:
                warnings.append(
                    f"pivot {pivot.name!r} on sheet {name!r} was built for "
                    f"columns {cached}, but {grain} now has {written}. "
                    f"Regenerate templates/workbook.xlsx with "
                    f"tools/make_workbook_template.py.")
    return warnings


def refresh_pivots_on_open(workbook: Workbook) -> int:
    """The cache in the template holds no records - the data is written after
    it. Excel rebuilds it from the table on open, which is what makes a
    pre-built pivot work against data it has never seen."""
    count = 0
    for name in workbook.sheetnames:
        for pivot in workbook[name]._pivots:
            pivot.cache.refreshOnLoad = True
            count += 1
    return count


def order_sheets(workbook: Workbook) -> None:
    """Known sheets first, in the documented order; the template's pivot
    sheets keep their relative order at the end."""
    known = [name for name in SHEET_ORDER if name in workbook.sheetnames]
    rest = [name for name in workbook.sheetnames if name not in known]
    workbook._sheets = [workbook[name] for name in known + rest]


CORE_PROPERTIES = "docProps/core.xml"
MODIFIED = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def normalise(path: Path, modified: dt.datetime) -> None:
    """Rewrite the zip so that the same model always produces the same bytes.

    Two things leak the clock into an xlsx. Every zip entry carries the time it
    was written, and openpyxl stamps docProps/core.xml with datetime.now() as
    it saves - after anything the caller set, so it cannot be prevented, only
    corrected here. Both are pinned to the model's own timestamp, and the
    reproducibility report_model.json guarantees then reaches the workbook
    (SPEC 5.2).
    """
    replacement = modified.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename))
                   for info in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, payload in entries:
            if info.filename == CORE_PROPERTIES:
                payload = MODIFIED.sub(rb"\g<1>" + replacement + rb"\g<2>", payload)
            copy = zipfile.ZipInfo(info.filename, date_time=ZIP_EPOCH)
            copy.compress_type = info.compress_type
            copy.external_attr = info.external_attr
            copy.internal_attr = info.internal_attr
            copy.create_system = info.create_system
            archive.writestr(copy, payload)


# --------------------------------------------------------------------------
# Driver
# --------------------------------------------------------------------------

def render(model: dict, out_dir: Path, lang: str, template: Path | None,
           model_dir: Path) -> tuple[Path, list[str]]:
    workbook, warnings = open_workbook(template)

    flat_columns = model.get("flat_columns") or {}
    flat = model.get("flat") or {}
    layouts = {name: sheet_columns(flat_columns.get(name, []), lang)
               for name in DATA_SHEETS}

    warnings += check_template(workbook, layouts)

    warnings += [f"{name}: no columns in the model, sheet left empty"
                 for name in DATA_SHEETS if not layouts[name]]

    missing = write_dashboard_sheet(sheet(workbook, "dashboard"), model, lang,
                                    model_dir)
    if missing:
        warnings.append(f"charts not yet rendered, dashboard has placeholders: "
                        f"{', '.join(missing)}")
    write_counts_sheet(sheet(workbook, "counts"), model, lang)
    for name in DATA_SHEETS:
        write_data_sheet(sheet(workbook, name), name, layouts[name],
                         flat.get(name) or [])
    write_definitions_sheet(sheet(workbook, "definitions"), model, lang)

    pivots = refresh_pivots_on_open(workbook)
    if not pivots:
        warnings.append("this workbook has no pre-built pivot tables")
    order_sheets(workbook)

    workbook.properties.creator = "mig pipeline"
    workbook.properties.title = t(model.get("report", {}), "title", lang, "")
    # No clock reading: the workbook is as reproducible as the model it came
    # from (SPEC 5.2).
    workbook.properties.created = stamp(model.get("generated_at"))
    workbook.properties.modified = workbook.properties.created

    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / OUTPUT_NAME
    workbook.save(out_path)
    normalise(out_path, workbook.properties.created)
    return out_path, warnings


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Render report_model.json to xlsx.")
    parser.add_argument("--model", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--lang", default="de", choices=("de", "en"))
    parser.add_argument("--template", default=DEFAULT_TEMPLATE, type=Path,
                        help="workbook holding the pre-built pivots "
                             "(default: templates/workbook.xlsx)")
    args = parser.parse_args(argv)

    if not args.model.is_file():
        print(f"--model: {args.model} not found", file=sys.stderr)
        return EXIT_TOOL
    try:
        model = json.loads(args.model.read_text(encoding="utf-8"))
    except ValueError as exc:
        print(f"--model: {args.model} is not valid JSON: {exc}", file=sys.stderr)
        return EXIT_TOOL

    try:
        out_path, warnings = render(model, args.out, args.lang, args.template,
                                    args.model.parent)
    except Exception as exc:  # pragma: no cover - surfaced as a tool error
        print(f"xlsx: {type(exc).__name__}: {exc}", file=sys.stderr)
        return EXIT_TOOL

    for warning in warnings:
        print(f"warning: {warning}", file=sys.stderr)
    counts = ", ".join(f"{name} {len(model.get('flat', {}).get(name) or [])}"
                       for name in DATA_SHEETS)
    print(f"xlsx: {out_path}")
    print(f"  {counts}")
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())

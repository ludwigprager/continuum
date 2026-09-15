"""Tests for tools/render/pptx.py and templates/deck.potx (M5).

The load-bearing ones:

  * test_the_committed_template_is_a_real_potx and
    test_python_pptx_alone_cannot_open_it - the corporate template (SPEC 12.2)
    will arrive as a .potx, and python-pptx refuses one outright. This pair is
    what says the loader handles the file that is actually coming, rather than
    a .pptx somebody renamed.
  * test_every_layout_the_renderer_names_exists_with_its_placeholders - SPEC
    6.4 asks for the layout mapping to live in one dict so swapping templates
    is a five-minute job. This is the test that makes a swap fail loudly at
    that dict instead of producing a plausible deck laid out wrongly.
  * test_no_row_is_dropped_when_a_table_is_split - a table too tall for one
    slide continues on the next. A row silently missing from a deck is how a
    meeting ends up being about the numbers.
"""
from __future__ import annotations

import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from conftest import renderer
from test_pipeline import FIXTURES, REPO, make_model, make_snapshot, run

PPTX = REPO / "tools" / "render" / "pptx.py"
CHARTS = REPO / "tools" / "render" / "charts.py"
TEMPLATE = REPO / "templates" / "deck.potx"

pptx_renderer = renderer("pptx")


def report_dir(tmp_path: Path, name: str = "report",
               projects: Path = FIXTURES) -> tuple[Path, dict]:
    """A report directory shaped like the one ./report.sh builds: the model
    next to the chart PNGs it names."""
    out = tmp_path / name
    out.mkdir(parents=True, exist_ok=True)
    snapshot = make_snapshot(tmp_path / f"{name}-tables", projects)
    model_path, model = make_model(tmp_path / f"{name}-tables", snapshot)
    (out / "report_model.json").write_bytes(model_path.read_bytes())
    run(CHARTS, "--model", str(out / "report_model.json"), "--out", str(out))
    return out, model


def render(out: Path, *extra: str) -> Path:
    run(PPTX, "--model", str(out / "report_model.json"), "--out", str(out),
        *extra)
    return out / "deck.pptx"


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    out, model = report_dir(tmp_path_factory.mktemp("pptx"))
    return Presentation(str(render(out))), model, out


def texts(slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()]


def tables(slide) -> list:
    return [shape.table for shape in slide.shapes if shape.has_table]


def cells(table) -> list[str]:
    return [table.cell(row, column).text
            for row in range(len(table.rows))
            for column in range(len(table.columns))]


def everything(presentation) -> str:
    """Every word in the deck, text frames and table cells alike.

    A table is a graphic frame, not a text frame, so a search that only walks
    `text_frame` misses every number on every table slide - which is most of
    the deck.
    """
    return "\n".join(
        [text for slide in presentation.slides for text in texts(slide)]
        + [cell for slide in presentation.slides
           for table in tables(slide) for cell in cells(table)])


# --------------------------------------------------------------------------
# The template: a .potx, which python-pptx refuses on its own
# --------------------------------------------------------------------------

def test_the_committed_template_is_a_real_potx():
    """Not a presentation with the extension changed. The corporate template
    (SPEC 12.2) will be the real thing, so the placeholder has to be too or
    the swap is the first time the loader meets one."""
    with zipfile.ZipFile(TEMPLATE) as archive:
        content_types = archive.read("[Content_Types].xml").decode()
    assert pptx_renderer.TEMPLATE_TYPE in content_types
    assert pptx_renderer.PRESENTATION_TYPE not in content_types


def test_python_pptx_alone_cannot_open_it():
    """The trap, stated as a test: without load_template this is where the
    deck renderer would stop, with an error that reads like a corrupt file."""
    with pytest.raises(ValueError) as excinfo:
        Presentation(io.BytesIO(TEMPLATE.read_bytes()))
    assert "not a PowerPoint file" in str(excinfo.value)


def test_the_loader_opens_it_anyway_without_touching_the_file():
    before = TEMPLATE.read_bytes()
    presentation, warnings = pptx_renderer.load_template(TEMPLATE)
    assert warnings == []
    assert len(presentation.slide_layouts) > 0
    assert TEMPLATE.read_bytes() == before


def test_every_layout_the_renderer_names_exists_with_its_placeholders():
    """SPEC 6.4: the layout/placeholder mapping is one dict so swapping the
    template is a five-minute job. Which only holds if the dict is checked."""
    presentation, _ = pptx_renderer.load_template(TEMPLATE)
    layouts = presentation.slide_layouts
    for role, mapping in pptx_renderer.LAYOUTS.items():
        assert mapping["layout"] < len(layouts), role
        layout = layouts[mapping["layout"]]
        assert layout.name == mapping["name"], role
        available = {ph.placeholder_format.idx for ph in layout.placeholders}
        wanted = {value for key, value in mapping.items()
                  if key not in ("layout", "name")}
        assert wanted <= available, f"{role}: missing {sorted(wanted - available)}"


def test_a_moved_layout_is_a_warning_not_a_wrong_deck(tmp_path):
    """A template swap that moves a layout must say so. A deck laid out on the
    wrong master looks plausible and is wrong, which is the worst pair of
    properties a report can have."""
    presentation, _ = pptx_renderer.load_template(TEMPLATE)
    presentation.slide_layouts[pptx_renderer.LAYOUTS["title"]["layout"]] \
        .name = "Something Else"
    warnings = pptx_renderer.check_layouts(presentation)
    assert any("Something Else" in warning for warning in warnings)


def test_a_missing_template_warns_loudly(tmp_path):
    out, _ = report_dir(tmp_path)
    proc = subprocess.run(
        [sys.executable, str(PPTX), "--model", str(out / "report_model.json"),
         "--out", str(out), "--template", str(tmp_path / "nope.potx")],
        capture_output=True, text=True, cwd=str(REPO))
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "no template" in proc.stderr
    assert (out / "deck.pptx").is_file()


# --------------------------------------------------------------------------
# The deck is the model, laid out
# --------------------------------------------------------------------------

def test_the_deck_opens_and_starts_with_the_report_title(deck):
    presentation, model, _ = deck
    assert len(presentation.slides) > 1
    assert model["report"]["title_de"] in "\n".join(texts(presentation.slides[0]))


def test_every_chart_gets_a_slide_with_the_picture_embedded(deck):
    """SPEC 2: charts are rendered once by charts.py and embedded, never
    redrawn here."""
    presentation, model, _ = deck
    assert model["charts"], "the fixtures should exercise at least one chart"
    pictures = [shape for slide in presentation.slides for shape in slide.shapes
                if getattr(shape, "image", None) is not None]
    assert len(pictures) == len(model["charts"])
    for picture, chart in zip(pictures, model["charts"]):
        assert picture.image.size == (chart["width_px"], chart["height_px"])


def test_the_chart_is_not_cropped(deck):
    """A picture placeholder crops an image to fill its frame, which cuts the
    category labels off a wide chart. The crop is undone and the picture
    refitted, so the chart in the deck is the chart in the PDF."""
    presentation, model, _ = deck
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "image", None) is None:
                continue
            assert (shape.crop_left, shape.crop_right,
                    shape.crop_top, shape.crop_bottom) == (0, 0, 0, 0)
            native = shape.image.size
            assert abs(shape.width / shape.height
                       - native[0] / native[1]) < 0.01, "aspect ratio changed"


def test_every_table_in_the_model_gets_a_slide(deck):
    presentation, model, _ = deck
    titles = {text for slide in presentation.slides for text in texts(slide)}
    for key, table in model["tables"].items():
        title = table["title_de"]
        assert any(title in text for text in titles), key


def test_an_unavailable_table_carries_the_models_note(deck):
    """n/a, not a crash and not a zero - and not a silently missing slide
    either (SPEC 6.3, 11)."""
    presentation, model, _ = deck
    unavailable = {key: table for key, table in model["tables"].items()
                   if not table.get("available")}
    assert unavailable, "the fixtures should name a field nobody has collected"
    words = everything(presentation)
    for key, table in unavailable.items():
        assert table["note_de"] in words, key


def test_unknown_is_never_dropped(deck):
    """SPEC 11: unknown is an explicit row in every distribution."""
    presentation, model, _ = deck
    everything = [cell for slide in presentation.slides
                  for table in tables(slide) for cell in cells(table)]
    for key, table in model["tables"].items():
        if not table.get("available") or table.get("kind") == "cross_tab":
            continue
        column = table["columns"][0]["key"]
        unknown = [row for row in table["rows"] if row[column] == "unknown"]
        assert unknown, f"{key} has no unknown row in the model"
        assert "unknown" in everything, key


def test_no_row_is_dropped_when_a_table_is_split(deck):
    """The appendix is longer than one slide holds. Every counting rule must
    still be in the deck (SPEC 7)."""
    presentation, model, _ = deck
    everything = [cell for slide in presentation.slides
                  for table in tables(slide) for cell in cells(table)]
    for definition in model["definitions"]:
        assert definition["key"] in everything, definition["key"]


def test_every_number_on_a_slide_is_the_models_number(deck):
    """The renderer lays out, it does not compute (SPEC 11)."""
    presentation, model, _ = deck
    everything = [cell for slide in presentation.slides
                  for table in tables(slide) for cell in cells(table)]
    for item in model["headline"]:
        assert item["label_de"] in everything, item["key"]
        assert str(item["value"]) in everything, item["key"]


def test_raw_and_verified_coverage_stay_two_numbers(deck):
    """SPEC 11: never merged into one, anywhere they appear."""
    presentation, model, _ = deck
    words = everything(presentation)
    coverage = model["coverage"]
    assert str(coverage["field_coverage_pct"]) in words
    assert str(coverage["verified_coverage_pct"]) in words
    for team in coverage["by_team"]:
        assert team["team_id"] in words


# --------------------------------------------------------------------------
# Reproducibility (SPEC 5.2)
# --------------------------------------------------------------------------

def test_the_deck_is_byte_identical_across_runs(tmp_path):
    out, _ = report_dir(tmp_path)
    first = render(out).read_bytes()
    (out / "deck.pptx").unlink()
    second = render(out).read_bytes()
    assert first == second


def test_the_timestamp_comes_from_the_model_not_the_clock(tmp_path):
    out, model = report_dir(tmp_path)
    first = render(out).read_bytes()
    model["generated_at"] = "2020-01-02T03:04:05+01:00"
    (out / "report_model.json").write_text(json.dumps(model), encoding="utf-8")
    second = render(out).read_bytes()
    assert first != second


def test_lang_switches_the_labels(tmp_path):
    """Primary audience is German, both languages must render (SPEC 1)."""
    out, model = report_dir(tmp_path)
    render(out, "--lang", "en")
    english = Presentation(str(out / "deck.pptx"))
    assert model["report"]["title_en"] in everything(english)


# --------------------------------------------------------------------------
# The moving target (SPEC 6.3)
# --------------------------------------------------------------------------

def test_a_chart_whose_png_is_missing_still_gets_its_slide(tmp_path):
    """charts.py runs first (SPEC 6.4). If it did not, the slide says which
    file is missing rather than disappearing out of the deck."""
    out, model = report_dir(tmp_path)
    (out / model["charts"][0]["path"]).unlink()
    proc = subprocess.run(
        [sys.executable, str(PPTX), "--model", str(out / "report_model.json"),
         "--out", str(out)], capture_output=True, text=True, cwd=str(REPO))
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert model["charts"][0]["key"] in proc.stderr

    presentation = Presentation(str(out / "deck.pptx"))
    assert model["charts"][0]["path"] in everything(presentation)


def test_a_missing_model_is_a_tool_error(tmp_path):
    proc = subprocess.run(
        [sys.executable, str(PPTX), "--model", str(tmp_path / "nope.json"),
         "--out", str(tmp_path)], capture_output=True, text=True, cwd=str(REPO))
    assert proc.returncode == 2
    assert "not found" in proc.stdout + proc.stderr

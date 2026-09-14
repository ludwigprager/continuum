"""Tests for tools/render/xlsx.py (M3).

The load-bearing ones:

  * test_the_prebuilt_pivot_survives_into_the_output - SPEC 6.4 warns that
    writing the workbook from scratch loses the pivots and the feature "fails
    silently". This is the test that makes it fail loudly instead.
  * test_counts_sheet_matches_the_model - the renderer lays out, it does not
    compute (SPEC 11). Every number on the sheet must be the model's number.
  * test_the_cross_tab_on_the_counts_sheet_matches_the_environments_sheet -
    the M3 acceptance criterion, minus the human: a pivot of site x OS family
    built from the environments sheet agrees with the counts sheet.
"""
from __future__ import annotations

import collections
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from test_pipeline import FIXTURES, REPO, make_model, make_snapshot, run

XLSX = REPO / "tools" / "render" / "xlsx.py"
TEMPLATE = REPO / "templates" / "workbook.xlsx"

DATA_SHEETS = ("projects", "datastores", "environments", "blockers",
               "dependencies")


def render(tmp_path: Path, *extra: str, projects: Path = FIXTURES) -> tuple[Path, dict]:
    """model -> workbook, the way report.sh does it."""
    snapshot = make_snapshot(tmp_path, projects)
    model_path, model = make_model(tmp_path, snapshot)
    out = tmp_path / "report"
    run(XLSX, "--model", str(model_path), "--out", str(out), *extra)
    return out / "report.xlsx", model


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    path, model = render(tmp_path_factory.mktemp("render"))
    return load_workbook(path), model, path


# --------------------------------------------------------------------------
# M3 acceptance (SPEC 10)
# --------------------------------------------------------------------------

def test_every_grain_is_a_real_excel_table(workbook):
    """"Open it and build a pivot without touching the data" needs a real
    table with a filter and a frozen header, not a range of cells."""
    wb, _, _ = workbook
    for name in DATA_SHEETS:
        ws = wb[name]
        assert list(ws.tables) == [f"tbl_{name}"], name
        assert ws.freeze_panes == "A2", name
        assert ws.tables[f"tbl_{name}"].autoFilter is not None, name


def test_the_prebuilt_pivot_survives_into_the_output(workbook):
    """SPEC 6.4: openpyxl cannot create a pivot table, so the template is
    loaded and written into. Write the workbook from scratch instead and the
    pivots are gone and the feature fails silently. Not silently."""
    wb, _, _ = workbook
    pivots = [pivot for name in wb.sheetnames for pivot in wb[name]._pivots]
    assert pivots, "the workbook lost the pre-built pivots from the template"
    for pivot in pivots:
        # The cache ships empty; Excel fills it from the sheet on open.
        assert pivot.cache.refreshOnLoad is True, pivot.name
        source = pivot.cache.cacheSource.worksheetSource
        assert source.name.startswith("tbl_"), "the pivot must read the table, "\
            "not a fixed cell range, or it stops at yesterday's row count"


def test_the_pivot_reads_the_columns_the_sheet_actually_has(workbook):
    """A schema change moves the columns. If the template is not regenerated
    the pivot counts the wrong field, so the two must be compared."""
    wb, _, _ = workbook
    for name in wb.sheetnames:
        for pivot in wb[name]._pivots:
            grain = pivot.cache.cacheSource.worksheetSource.name[len("tbl_"):]
            headers = [cell.value for cell in wb[grain][1]]
            assert [f.name for f in pivot.cache.cacheFields] == headers


def test_the_cross_tab_on_the_counts_sheet_matches_the_environments_sheet(workbook):
    """The acceptance criterion without the human in the loop: count site x OS
    family off the environments sheet and compare it to the counts sheet."""
    wb, model, _ = workbook
    ws = wb["environments"]
    headers = [cell.value for cell in ws[1]]
    site, os_family = headers.index("datacenter_code"), headers.index("os_family_code")
    tallied: collections.Counter = collections.Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        tallied[(row[site], row[os_family])] += 1

    table = model["tables"]["site_x_os"]
    row_key = table["columns"][0]["key"]
    for record in table["rows"]:
        for column in table["columns"][1:]:
            if column["key"] == "total":
                continue
            assert record[column["key"]] == tallied[(record[row_key], column["key"])], \
                f"{record[row_key]} x {column['key']}"
    assert sum(tallied.values()) == sum(r["total"] for r in table["rows"])


# --------------------------------------------------------------------------
# The renderer computes nothing (SPEC 11)
# --------------------------------------------------------------------------

def read_counts(ws) -> dict[str, dict[str, int]]:
    """The counts sheet, back into {table key: {code: value}}."""
    blocks: dict[str, dict[str, int]] = {}
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for index, row in enumerate(rows):
        marker = str(row[0] or "")
        if "  |  " not in marker:
            continue
        key = marker.split("  |  ")[0]
        header = rows[index + 1]
        if header[0] != "Code":
            continue           # a cross-tab, checked separately
        block = {}
        for record in rows[index + 2:]:
            if not record[0]:
                break
            block[str(record[0])] = record[2]
        blocks[key] = block
    return blocks


def test_counts_sheet_matches_the_model(workbook):
    wb, model, _ = workbook
    blocks = read_counts(wb["counts"])
    assert blocks, "the counts sheet has no frequency tables"
    for key, block in blocks.items():
        table = model["tables"][key]
        column = table["columns"][0]["key"]
        expected = {str(r[column]): r["value"] for r in table["rows"]}
        assert block == expected, key


def test_unknown_is_on_the_sheet_never_dropped(workbook):
    """SPEC 11. The unknown bar is the honest one, especially early on."""
    wb, model, _ = workbook
    for key, block in read_counts(wb["counts"]).items():
        assert "unknown" in block, key


def test_a_field_nobody_collected_yet_is_printed_as_na_not_omitted(workbook):
    """SPEC 6.3: a field named in the spec but absent from the data renders
    n/a. Dropping the table would hide that somebody asked for it."""
    wb, model, _ = workbook
    absent = [key for key, table in model["tables"].items()
              if not table.get("available")]
    assert absent, "the fixtures no longer cover this case"
    text = "\n".join(str(cell.value) for row in wb["counts"].iter_rows()
                     for cell in row if cell.value)
    for key in absent:
        assert key in text, f"{key} vanished from the counts sheet"
        assert model["tables"][key]["note_de"] in text


def test_coded_fields_get_a_code_column_and_a_label_column(workbook):
    """SPEC 6.4: pivot on the code so the sort order is right, display the
    label. Both come from the model; the renderer never sees a taxonomy."""
    wb, model, _ = workbook
    for name in DATA_SHEETS:
        headers = [cell.value for cell in wb[name][1]]
        coded = [c["key"] for c in model["flat_columns"][name] if c["coded"]]
        assert coded or name == "dependencies", name
        for key in coded:
            assert f"{key}_code" in headers and f"{key}_label" in headers, key

    headers = [cell.value for cell in wb["projects"][1]]
    code_at = headers.index("migration_status_code")
    label_at = headers.index("migration_status_label")
    rows = {r[0]: r for r in wb["projects"].iter_rows(min_row=2, values_only=True)}
    by_id = {r["project_id"]: r for r in model["flat"]["projects"]}
    for project_id, row in rows.items():
        expected = by_id[project_id]
        assert row[code_at] == expected["migration_status"]
        assert row[label_at] == expected["migration_status_label_de"]


def test_definitions_are_printed_next_to_the_numbers(workbook):
    """SPEC 7: a number and the rule it was counted under never travel apart."""
    wb, model, _ = workbook
    keys = [row[0] for row in wb["definitions"].iter_rows(min_row=2, values_only=True)]
    assert keys == [d["key"] for d in model["definitions"]]
    texts = [row[3] for row in wb["definitions"].iter_rows(min_row=2, values_only=True)]
    assert all(texts), "every definition needs its text on the sheet"


def test_the_dashboard_carries_raw_and_verified_coverage_separately(workbook):
    """SPEC 11: two numbers, everywhere they appear."""
    wb, model, _ = workbook
    values = {}
    for row in wb["dashboard"].iter_rows(values_only=True):
        if row[0] and row[1] is not None:
            values[str(row[0])] = row[1]
    assert values["field_coverage_pct"] == model["coverage"]["field_coverage_pct"]
    assert values["verified_coverage_pct"] == model["coverage"]["verified_coverage_pct"]
    assert values["git_sha"] == model["provenance"]["git_sha"]


# --------------------------------------------------------------------------
# Traps
# --------------------------------------------------------------------------

def test_a_missing_template_is_loud(tmp_path):
    """The one failure mode SPEC 6.4 calls out by name. It must be impossible
    to ship a pivot-less workbook without being told."""
    snapshot = make_snapshot(tmp_path)
    model_path, _ = make_model(tmp_path, snapshot)
    out = tmp_path / "no-template"
    proc = subprocess.run(
        [sys.executable, str(XLSX), "--model", str(model_path), "--out", str(out),
         "--template", str(tmp_path / "nope.xlsx")],
        capture_output=True, text=True, cwd=str(REPO))
    assert proc.returncode == 0, proc.stderr
    assert "no template" in proc.stderr
    assert "pivot" in proc.stderr.lower()
    wb = load_workbook(out / "report.xlsx")
    assert not [p for name in wb.sheetnames for p in wb[name]._pivots]


def test_a_value_that_looks_like_a_formula_stays_text(tmp_path):
    """Several hundred people hand-edit these files. A name beginning with an
    equals sign is a typo, not something the report should evaluate."""
    projects = tmp_path / "projects"
    projects.mkdir()
    (projects / "odd.yaml").write_text(
        'schema_version: 1\nid: odd\nmy_mandatory_field: "=1+1"\n'
        'name: "=cmd|calc"\n', encoding="utf-8")
    path, _ = render(tmp_path, projects=projects)
    ws = load_workbook(path)["projects"]
    headers = [cell.value for cell in ws[1]]
    row = next(ws.iter_rows(min_row=2, max_row=2))
    for column in ("name", "my_mandatory_field"):
        cell = row[headers.index(column)]
        assert cell.data_type == "s", f"{column} became a formula"
        assert str(cell.value).startswith("=")


def test_the_workbook_is_byte_identical_across_runs(tmp_path):
    """The model is reproducible (SPEC 5.2); the workbook built from it has no
    excuse not to be. Without pinning the zip timestamps it never is.

    One model, rendered twice: the model records which snapshot directory
    produced it, so building two of them would differ for a reason that has
    nothing to do with the renderer.
    """
    snapshot = make_snapshot(tmp_path)
    model_path, _ = make_model(tmp_path, snapshot)
    written = []
    for name in ("a", "b"):
        out = tmp_path / name
        run(XLSX, "--model", str(model_path), "--out", str(out))
        written.append((out / "report.xlsx").read_bytes())
    assert written[0] == written[1]


def test_an_empty_grain_still_produces_a_filterable_sheet(tmp_path):
    """The first day nobody records a blocker, the blockers sheet must still
    be a table someone can pivot - an Excel table needs a body row."""
    projects = tmp_path / "projects"
    projects.mkdir()
    (projects / "bare.yaml").write_text(
        "schema_version: 1\nid: bare\nmy_mandatory_field: TODO\n", encoding="utf-8")
    path, model = render(tmp_path, projects=projects)
    assert model["flat"]["blockers"] == []
    ws = load_workbook(path)["blockers"]
    assert list(ws.tables) == ["tbl_blockers"]
    assert ws.tables["tbl_blockers"].ref.endswith("2")
    assert [cell.value for cell in ws[1]], "the header must survive an empty table"


def test_the_template_is_committed():
    """It is not generated by the daily run: if it goes missing, every report
    from then on quietly loses its pivots."""
    assert TEMPLATE.is_file(), (
        "templates/workbook.xlsx is missing. Rebuild it with "
        "./shell.sh python3 tools/make_workbook_template.py "
        "--model out/reports/<date>/report_model.json")


def test_lang_switches_the_labels(tmp_path):
    """Primary audience is German, but the labels must support both (SPEC 1)."""
    snapshot = make_snapshot(tmp_path)
    model_path, model = make_model(tmp_path, snapshot)
    sheets = {}
    for lang in ("de", "en"):
        out = tmp_path / lang
        run(XLSX, "--model", str(model_path), "--out", str(out), "--lang", lang)
        ws = load_workbook(out / "report.xlsx")["projects"]
        headers = [cell.value for cell in ws[1]]
        column = headers.index("migration_status_label")
        sheets[lang] = {row[0]: row[column]
                        for row in ws.iter_rows(min_row=2, values_only=True)}
    by_id = {r["project_id"]: r for r in model["flat"]["projects"]}
    for lang in ("de", "en"):
        for project_id, label in sheets[lang].items():
            assert label == by_id[project_id][f"migration_status_label_{lang}"]
    assert sheets["de"] != sheets["en"], "the fixtures should differ in both"
